#!/usr/bin/env python3
"""
iata_to_excel.py
=================
Converts an IATA / BSP "FCAGBILLDET" (Agent Billing Details) PDF into a
formatted Excel workbook, with one row per ticket / transaction document.

USAGE (from cmd / terminal):
    python iata_to_excel.py "EG_FCAGBILLDET_9020124_20260802.PDF"
    python iata_to_excel.py "EG_FCAGBILLDET_9020124_20260802.PDF" "output.xlsx"

If no output path is given, the script writes "<input_name>.xlsx" next to
the input file.

Requires:
    pip install pdfplumber openpyxl
"""

import sys
import re
import os
from collections import defaultdict

try:
    import pdfplumber
except ImportError:
    sys.exit("Missing dependency 'pdfplumber'. Install it with:\n"
              "    pip install pdfplumber openpyxl")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Missing dependency 'openpyxl'. Install it with:\n"
              "    pip install pdfplumber openpyxl")


# --------------------------------------------------------------------------
# Column layout (matches the "SCOPE / COMBINED" detail table header row):
#   AIR TRNC Number Date CPUI Code STAT FOP Amount Amount TAX F&C PEN
#   Amount Rate Amt Rate Amt Comm Payable
# --------------------------------------------------------------------------
FIELD_NAMES = [
    "AIR", "TRNC", "Document Number", "Issue Date", "CPUI", "NR Code",
    "STAT", "FOP", "Transaction Amount", "FARE Amount", "TAX", "F&C",
    "PEN", "COBL Amount", "STD Comm Rate", "STD Comm Amt",
    "SUPP Comm Rate", "SUPP Comm Amt", "Tax on Comm", "Balance Payable",
]

AMOUNT_CODE_RE = re.compile(r"^(-?[\d,]+\.\d{2})\s*\*?([A-Z0-9]{1,3})$")
PLAIN_AMOUNT_RE = re.compile(r"^-?[\d,]+\.\d{2}\*?$")
AIR_RE = re.compile(r"^\d{2,3}[A-Z]?$")

# Lines that are section/grand totals, banners, or page furniture rather
# than an actual ticket/document row. These must never be merged into a
# transaction row, and must force-flush whatever row is currently open.
BANNER_PREFIXES = (
    "ISSUES TOTAL", "REFUNDS TOTAL", "DEBIT MEMOS TOTAL", "CREDIT MEMOS TOTAL",
    "COMBINED TOTALS", "BSP TOTALS", "WEBSALES-EDIS TOTALS",
    "ISSUES CA", "REFUNDS CA", "DEBIT MEMOS CA", "CREDIT MEMOS",
    "GRAND TOTAL", "EGYPT", "BSP TOTAL", "WEBSALES-EDIS TOTAL",
    "CATEGORY",
)


def cluster_words_by_line(words, tolerance=2.5):
    """Group words into visual lines based on their 'top' coordinate."""
    words = sorted(words, key=lambda w: w["top"])
    clusters = []
    cur, cur_top = [], None
    for w in words:
        if cur_top is None or abs(w["top"] - cur_top) <= tolerance:
            cur.append(w)
            cur_top = w["top"] if cur_top is None else cur_top
        else:
            clusters.append(cur)
            cur, cur_top = [w], w["top"]
    if cur:
        clusters.append(cur)
    return clusters


def build_column_bounds(header_words):
    """
    Given the 20 header words of the detail table (already sorted by x0),
    return a list of (name, x_start, x_end) boundaries using midpoints
    between consecutive header word positions.
    """
    xs = [w["x0"] for w in header_words]
    bounds = []
    for i, name in enumerate(FIELD_NAMES):
        start = -1 if i == 0 else (xs[i - 1] + xs[i]) / 2
        end = 99999 if i == len(FIELD_NAMES) - 1 else (xs[i] + xs[i + 1]) / 2
        bounds.append((name, start, end))
    return bounds


def assign_to_column(word, bounds):
    x0 = word["x0"]
    for name, start, end in bounds:
        if start <= x0 < end:
            return name
    return None


def parse_amount_code(token):
    """Split a merged 'amount+code' token like '214.00E3' into (amount, code)."""
    m = AMOUNT_CODE_RE.match(token)
    if m:
        return m.group(1), m.group(2)
    if PLAIN_AMOUNT_RE.match(token):
        return token, ""
    return None, token


def to_number(s):
    if s is None or s == "":
        return None
    s = s.replace(",", "").replace("*", "")
    try:
        return float(s)
    except ValueError:
        return None


def is_main_row_start(cluster, bounds):
    """A main transaction row starts with an AIR-code-like token in the
    AIR column, together with a TRNC-like word in the TRNC column."""
    air_word = None
    trnc_word = None
    for w in cluster:
        col = assign_to_column(w, bounds)
        if col == "AIR" and air_word is None:
            air_word = w["text"]
        if col == "TRNC" and trnc_word is None:
            trnc_word = w["text"]
    if air_word and AIR_RE.match(air_word) and trnc_word:
        return True
    return False


def parse_detail_pages(pdf, category_default="BSP"):
    """
    Walk every page of the PDF, find every 'SCOPE / COMBINED' detail table,
    and return a list of row-dicts (one per ticket/document), tagged with
    Category (BSP / WEBSALES-EDIS) and Section (ISSUES / REFUNDS /
    DEBIT MEMOS / CREDIT MEMOS).
    """
    rows = []
    current_category = category_default
    current_section = "ISSUES"
    current_row = None
    bounds = None

    for page in pdf.pages:
        text = page.extract_text() or ""

        # Update category / section context from banner lines on this page.
        if "WEBSALES-EDIS" in text:
            current_category = "WEBSALES-EDIS"
        if re.search(r"CATEGORY\s*\n?\s*BSP\b", text) or (
            "CATEGORY" in text and "BSP" in text and "WEBSALES" not in text
        ):
            current_category = "BSP"
        if "*** ISSUES" in text:
            current_section = "ISSUES"
        if "*** REFUNDS" in text:
            current_section = "REFUNDS"
        if "*** DEBIT MEMOS" in text:
            current_section = "DEBIT MEMOS"
        if "*** CREDIT MEMOS" in text:
            current_section = "CREDIT MEMOS"

        words = page.extract_words()
        if not words:
            continue
        clusters = cluster_words_by_line(words)

        # Find this page's header row to (re)build column bounds.
        header_cluster = None
        for c in clusters:
            texts = [w["text"] for w in c]
            if "AIR" in texts and "TRNC" in texts and "Number" in texts:
                header_cluster = c
                break
        if header_cluster:
            header_cluster = sorted(header_cluster, key=lambda w: w["x0"])
            if len(header_cluster) == len(FIELD_NAMES):
                bounds = build_column_bounds(header_cluster)

        if bounds is None:
            continue  # no detail table encountered yet

        for c in clusters:
            texts = " ".join(w["text"] for w in c)

            # Section / category marker lines - already handled via page text,
            # but re-check per-line too (some pages mix sections).
            if texts.startswith("*** ISSUES"):
                if current_row is not None:
                    rows.append(current_row)
                    current_row = None
                current_section = "ISSUES"
                continue
            if texts.startswith("*** REFUNDS"):
                if current_row is not None:
                    rows.append(current_row)
                    current_row = None
                current_section = "REFUNDS"
                continue
            if texts.startswith("*** DEBIT MEMOS"):
                if current_row is not None:
                    rows.append(current_row)
                    current_row = None
                current_section = "DEBIT MEMOS"
                continue
            if texts.startswith("*** CREDIT MEMOS"):
                if current_row is not None:
                    rows.append(current_row)
                    current_row = None
                current_section = "CREDIT MEMOS"
                continue
            if texts.startswith(("SCOPE", "AIR TRNC", "Document Issue",
                                  "FCAGBILLDET", "Billing Period")) or re.match(
                    r"^\d{1,2}-[A-Z]{3}\b", texts):
                continue

            # Grand-total / section-total / banner lines: flush and skip.
            if texts.startswith(BANNER_PREFIXES):
                if current_row is not None:
                    rows.append(current_row)
                    current_row = None
                continue

            # +RTDN continuation line -> attach to the last main row.
            if any(w["text"] == "+RTDN:" for w in c):
                if current_row is not None:
                    nums = [w["text"] for w in c if w["text"] not in ("+RTDN:", "EX")]
                    if nums:
                        current_row["RTDN"] = nums[0]
                        if len(nums) > 1:
                            current_row["RTDN Code"] = nums[1]
                continue

            if is_main_row_start(c, bounds):
                # flush previous row
                if current_row is not None:
                    rows.append(current_row)
                current_row = {name: "" for name in FIELD_NAMES}
                current_row["Category"] = current_category
                current_row["Section"] = current_section
                current_row["RTDN"] = ""
                current_row["RTDN Code"] = ""
                current_row["PEN Code"] = ""
                current_row["Tax Breakdown"] = []
                current_row["F&C Breakdown"] = []

            if current_row is None:
                continue  # stray line before any real row (e.g. totals)

            # Distribute this cluster's words into columns.
            for w in c:
                col = assign_to_column(w, bounds)
                if col is None:
                    continue
                token = w["text"]
                if col in ("TAX", "F&C"):
                    amt, code = parse_amount_code(token)
                    if amt is not None:
                        target = "Tax Breakdown" if col == "TAX" else "F&C Breakdown"
                        current_row[target].append((amt, code))
                    continue
                if col == "PEN":
                    amt, code = parse_amount_code(token)
                    if amt is not None:
                        if current_row.get("PEN", "") == "":
                            current_row["PEN"] = amt
                        if code:
                            current_row["PEN Code"] = code
                    continue
                # Plain columns: keep first non-empty value seen
                # (rate/amount columns only ever appear on the main line).
                if current_row.get(col, "") == "":
                    current_row[col] = token
                else:
                    # AIR/TRNC etc. shouldn't repeat, but guard anyway
                    pass

        # end of clusters loop for this page

    if current_row is not None:
        rows.append(current_row)

    return rows


def finalize_rows(raw_rows):
    """Turn the raw parsed rows into flat dicts ready for Excel, computing
    TAX / F&C totals from their breakdown lists."""
    final = []
    for r in raw_rows:
        tax_total = sum(to_number(a) or 0 for a, _ in r["Tax Breakdown"])
        fc_total = sum(to_number(a) or 0 for a, _ in r["F&C Breakdown"])
        tax_codes = ", ".join(f"{a} {c}".strip() for a, c in r["Tax Breakdown"])
        fc_codes = ", ".join(f"{a} {c}".strip() for a, c in r["F&C Breakdown"])

        out = {
            "Category": r["Category"],
            "Section": r["Section"],
            "AIR": r["AIR"],
            "TRNC": r["TRNC"],
            "Document Number": r["Document Number"],
            "Issue Date": r["Issue Date"],
            "CPUI": r["CPUI"],
            "NR Code": r["NR Code"],
            "STAT": r["STAT"],
            "FOP": r["FOP"],
            "Transaction Amount": to_number(r["Transaction Amount"]),
            "FARE Amount": to_number(r["FARE Amount"]),
            "TAX Total": round(tax_total, 2) if r["Tax Breakdown"] else to_number(r["TAX"]),
            "TAX Breakdown": tax_codes,
            "F&C Total": round(fc_total, 2) if r["F&C Breakdown"] else to_number(r["F&C"]),
            "F&C Breakdown": fc_codes,
            "PEN": to_number(r["PEN"]),
            "PEN Code": r.get("PEN Code", ""),
            "COBL Amount": to_number(r["COBL Amount"]),
            "STD Comm Rate": to_number(r["STD Comm Rate"]),
            "STD Comm Amt": to_number(r["STD Comm Amt"]),
            "SUPP Comm Rate": to_number(r["SUPP Comm Rate"]),
            "SUPP Comm Amt": to_number(r["SUPP Comm Amt"]),
            "Tax on Comm": to_number(r["Tax on Comm"]),
            "Balance Payable": to_number(r["Balance Payable"]),
            "Related Doc (RTDN)": r["RTDN"],
            "RTDN Code": r["RTDN Code"],
        }
        final.append(out)
    return final


def parse_summary_table(pdf):
    """Parse the front-page SUMMARY table (BSP TOTAL / WEBSALES-EDIS TOTAL /
    GRAND TOTAL)."""
    summary_rows = []
    page = pdf.pages[0]
    text = page.extract_text() or ""
    lines = text.splitlines()
    headers = ["Category", "Transaction Amount", "FARE Amount", "TAX",
               "F&C", "PEN", "COBL Amount", "STD Comm Amt", "SUPP Comm Amt",
               "Tax on Comm", "Balance Payable"]
    for line in lines:
        m = re.match(
            r"^(BSP TOTAL|WEBSALES-EDIS TOTAL|GRAND TOTAL \(EGP\))\s+"
            r"([\d,\.\-]+)\s+([\d,\.\-]+)\s+([\d,\.\-]+)\s+([\d,\.\-]+)\s+"
            r"([\d,\.\-]+)\s+([\d,\.\-]+)\s+([\d,\.\-]+)\s+([\d,\.\-]+)\s+"
            r"([\d,\.\-]+)\s+([\d,\.\-]+)\s*$",
            line.strip(),
        )
        if m:
            vals = [m.group(1)] + [to_number(x) for x in m.groups()[1:]]
            summary_rows.append(dict(zip(headers, vals)))
    return summary_rows


def autosize_and_style(ws, n_header_cols):
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    for col_idx in range(1, n_header_cols + 1):
        letter = get_column_letter(col_idx)
        max_len = max(
            (len(str(cell.value)) for cell in ws[letter] if cell.value is not None),
            default=8,
        )
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 45)


def write_excel(detail_rows, summary_rows, out_path):
    """Write ALL transactions to ONE sheet.
    `out_path` can be a file path or a file-like object (e.g. BytesIO).
    `summary_rows` is kept in the signature for compatibility but not written."""
    wb = Workbook()
    ws = wb.active
    ws.title = "IATA"
    if detail_rows:
        headers = list(detail_rows[0].keys())
        ws.append(headers)
        for r in detail_rows:
            ws.append([r.get(h, "") for h in headers])
        autosize_and_style(ws, len(headers))
    wb.save(out_path)


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    in_path = sys.argv[1]
    if not os.path.isfile(in_path):
        sys.exit(f"File not found: {in_path}")

    if len(sys.argv) >= 3:
        out_path = sys.argv[2]
    else:
        base, _ = os.path.splitext(in_path)
        out_path = base + ".xlsx"

    print(f"Reading: {in_path}")
    with pdfplumber.open(in_path) as pdf:
        summary_rows = parse_summary_table(pdf)
        raw_rows = parse_detail_pages(pdf)

    detail_rows = finalize_rows(raw_rows)
    print(f"Parsed {len(detail_rows)} transaction rows.")

    write_excel(detail_rows, summary_rows, out_path)
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    main()
